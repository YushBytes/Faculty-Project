"""Import pipeline through the API: stage -> preview -> fix/exclude/map -> confirm."""

import io
from datetime import UTC, datetime, timedelta
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.modules.assessments.models import AssessmentResult, ResultSource, ResultStatus
from app.modules.audit.models import AuditLog
from app.modules.imports.models import ImportBatch
from app.modules.users.models import Role, User
from tests.conftest import auth_headers


def xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def setup(cse, cse_offering, make_student, make_assessment):
    students = [make_student(cse, enroll_in=[cse_offering]) for _ in range(3)]
    return {
        "offering": cse_offering,
        "students": students,
        "ct1": make_assessment(cse_offering, "CT1", max_marks="50"),
        "ct2": make_assessment(cse_offering, "CT2", max_marks="50"),
    }


def upload(client, user, setup, rows, name="marks.xlsx"):
    content = xlsx(rows) if name.endswith(".xlsx") else rows
    return client.post(
        f"/api/v1/offerings/{setup['offering'].id}/imports",
        files={"file": (name, content)},
        headers=auth_headers(user),
    )


def result_count(db: Session) -> int:
    return db.scalar(select(func.count()).select_from(AssessmentResult))


def sheet(setup, ct1_values, ct2_values=None):
    s = setup["students"]
    rows = [["Register No", "Name", "CT1 (max 50)", "CT2"]]
    for i, student in enumerate(s):
        ct2 = ct2_values[i] if ct2_values else "30"
        rows.append([student.register_number, student.full_name, ct1_values[i], ct2])
    return rows


class TestStageAndConfirm:
    def test_full_flow_writes_nothing_until_confirm(
        self, client: TestClient, faculty: User, setup, db_session: Session
    ) -> None:
        response = upload(client, faculty, setup, sheet(setup, [45, None, "AB"]))
        assert response.status_code == 201, response.text
        preview = response.json()
        assert preview["batch"]["status"] == "previewed"
        assert preview["summary"]["errors"] == 0
        assert preview["summary"]["will_create"] == 6
        assert preview["summary"]["blank_cells"] == 1
        assert result_count(db_session) == 0  # staged only

        batch_id = preview["batch"]["id"]
        confirmed = client.post(
            f"/api/v1/imports/{batch_id}/confirm", headers=auth_headers(faculty)
        )
        assert confirmed.status_code == 200, confirmed.text
        assert (confirmed.json()["created"], confirmed.json()["updated"]) == (6, 0)

        rows = {
            (r.student_id, r.assessment_id): r for r in db_session.scalars(select(AssessmentResult))
        }
        s, ct1 = setup["students"], setup["ct1"]
        blank = rows[(s[1].id, ct1.id)]
        assert (blank.status, blank.score) == (ResultStatus.ABSENT, None)  # never 0
        assert rows[(s[0].id, ct1.id)].score == Decimal("45.00")
        assert rows[(s[2].id, ct1.id)].status is ResultStatus.ABSENT
        assert all(r.source is ResultSource.IMPORT for r in rows.values())
        assert all(str(r.import_batch_id) == batch_id for r in rows.values())

        batch = client.get(f"/api/v1/imports/{batch_id}", headers=auth_headers(faculty)).json()
        assert batch["status"] == "committed" and batch["committed_at"]
        again = client.post(f"/api/v1/imports/{batch_id}/confirm", headers=auth_headers(faculty))
        assert again.status_code == 409

    def test_blocking_errors_prevent_confirm(
        self, client: TestClient, faculty: User, setup, db_session: Session
    ) -> None:
        preview = upload(client, faculty, setup, sheet(setup, [45, 99, "twelve"])).json()
        assert preview["summary"]["errors"] == 2
        assert preview["summary"]["exceeds_max"] == 1 and preview["summary"]["invalid_values"] == 1
        response = client.post(
            f"/api/v1/imports/{preview['batch']['id']}/confirm", headers=auth_headers(faculty)
        )
        assert response.status_code == 422
        details = response.json()["error"]["details"]
        assert {d["code"] for d in details} == {"score_above_max", "invalid_number"}
        assert {tuple(d["loc"][:2]) for d in details} == {("row", 3), ("row", 4)}
        assert result_count(db_session) == 0

    def test_fix_and_exclude_then_confirm(
        self, client: TestClient, faculty: User, setup, db_session: Session
    ) -> None:
        preview = upload(client, faculty, setup, sheet(setup, [45, 99, None])).json()
        batch_id = preview["batch"]["id"]
        url = f"/api/v1/imports/{batch_id}"
        headers = auth_headers(faculty)

        fixed = client.post(
            f"{url}/fix",
            json={
                "fixes": [
                    {"row": 3, "column": "CT1 (max 50)", "value": "49"},
                    {"row": 4, "column": "CT1 (max 50)", "value": "0"},
                ]
            },
            headers=headers,
        ).json()
        assert fixed["summary"]["errors"] == 0 and fixed["summary"]["blank_cells"] == 0
        row4 = next(r for r in fixed["rows"] if r["row"] == 4)
        zero = next(c for c in row4["cells"] if c["column"] == "CT1 (max 50)")
        assert (zero["raw"], zero["value"], zero["fixed"], zero["score"]) == (None, "0", True, 0.0)

        excluded = client.post(f"{url}/exclude", json={"rows": [2]}, headers=headers).json()
        assert excluded["summary"]["excluded_rows"] == 1
        assert client.post(f"{url}/confirm", headers=headers).json()["created"] == 4

        # The blank -> 0 override is in the audit log.
        fixes = db_session.scalars(
            select(AuditLog).where(AuditLog.entity == "import_batch", AuditLog.action == "fix")
        ).all()
        blank_to_zero = next(f for f in fixes if f.new_value["row"] == 4)
        assert (
            blank_to_zero.old_value["uploaded"] is None and blank_to_zero.new_value["value"] == "0"
        )
        assert blank_to_zero.actor_id == faculty.id

    def test_reset_fix(self, client: TestClient, faculty: User, setup) -> None:
        preview = upload(client, faculty, setup, sheet(setup, [45, 99, 10])).json()
        url = f"/api/v1/imports/{preview['batch']['id']}/fix"
        headers = auth_headers(faculty)
        client.post(
            url,
            json={"fixes": [{"row": 3, "column": "CT1 (max 50)", "value": "9"}]},
            headers=headers,
        )
        reset = client.post(
            url,
            json={"fixes": [{"row": 3, "column": "CT1 (max 50)", "reset": True}]},
            headers=headers,
        ).json()
        assert reset["summary"]["errors"] == 1

    def test_fix_rejects_unknown_row_or_column(
        self, client: TestClient, faculty: User, setup
    ) -> None:
        preview = upload(client, faculty, setup, sheet(setup, [1, 2, 3])).json()
        url = f"/api/v1/imports/{preview['batch']['id']}/fix"
        bad = client.post(
            url,
            json={
                "fixes": [
                    {"row": 99, "column": "CT1 (max 50)", "value": "1"},
                    {"row": 2, "column": "Nope", "value": "1"},
                ]
            },
            headers=auth_headers(faculty),
        )
        assert bad.status_code == 422 and len(bad.json()["error"]["details"]) == 2

    def test_column_mapping(self, client: TestClient, faculty: User, setup) -> None:
        rows = [["Reg No", "Cycle Test II"]] + [[s.register_number, 20] for s in setup["students"]]
        preview = upload(client, faculty, setup, rows).json()
        assert preview["summary"]["errors"] == 1
        column = next(c for c in preview["columns"] if c["header"] == "Cycle Test II")
        assert column["issues"][0]["code"] == "missing_assessment_definition"

        mapped = client.post(
            f"/api/v1/imports/{preview['batch']['id']}/mapping",
            json={"mappings": {"Cycle Test II": str(setup["ct2"].id)}},
            headers=auth_headers(faculty),
        ).json()
        assert mapped["summary"]["errors"] == 0 and mapped["summary"]["will_create"] == 3

    def test_overwrite_is_audited_with_batch(
        self, client: TestClient, faculty: User, setup, db_session: Session
    ) -> None:
        first = upload(client, faculty, setup, sheet(setup, [45, 40, 30])).json()
        client.post(
            f"/api/v1/imports/{first['batch']['id']}/confirm", headers=auth_headers(faculty)
        )
        second = upload(client, faculty, setup, sheet(setup, [46, 40, 30])).json()
        s = second["summary"]
        assert (s["will_update"], s["unchanged"]) == (1, 5)
        codes = {i["code"] for i in second["file_issues"]}
        assert "will_overwrite" in codes
        confirmed = client.post(
            f"/api/v1/imports/{second['batch']['id']}/confirm", headers=auth_headers(faculty)
        ).json()
        assert confirmed["updated"] == 1
        log = db_session.scalar(
            select(AuditLog).where(
                AuditLog.entity == "assessment_result", AuditLog.action == "update"
            )
        )
        assert log.old_value["score"] == "45.00" and log.new_value["score"] == "46.00"
        assert log.new_value["import_batch_id"] == second["batch"]["id"]

    def test_same_file_twice_warns(self, client: TestClient, faculty: User, setup) -> None:
        content = xlsx(sheet(setup, [45, 40, 30]))
        url = f"/api/v1/offerings/{setup['offering'].id}/imports"
        headers = auth_headers(faculty)
        first = client.post(url, files={"file": ("m.xlsx", content)}, headers=headers).json()
        assert "duplicate_file" not in {i["code"] for i in first["file_issues"]}
        client.post(f"/api/v1/imports/{first['batch']['id']}/confirm", headers=headers)

        again = client.post(url, files={"file": ("copy.xlsx", content)}, headers=headers).json()
        warning = next(i for i in again["file_issues"] if i["code"] == "duplicate_file")
        assert first["batch"]["id"] in warning["message"]
        assert again["summary"]["unchanged"] == 6

    def test_nothing_to_import(self, client: TestClient, faculty: User, setup) -> None:
        rows = sheet(setup, [45, 40, 30])
        first = upload(client, faculty, setup, rows).json()
        client.post(
            f"/api/v1/imports/{first['batch']['id']}/confirm", headers=auth_headers(faculty)
        )
        second = upload(client, faculty, setup, rows).json()
        response = client.post(
            f"/api/v1/imports/{second['batch']['id']}/confirm", headers=auth_headers(faculty)
        )
        assert (
            response.status_code == 422
            and "Nothing to import" in response.json()["error"]["message"]
        )


class TestFormats:
    def test_long_csv_for_single_assessment(
        self, client: TestClient, faculty: User, setup, db_session: Session
    ) -> None:
        s = setup["students"]
        csv = (
            "Register No,Assessment,Score,Max Score,Percentage\n"
            f"{s[0].register_number},CT1,40,50,80\n"
            f"{s[1].register_number},CT1,AB,50,\n"
        )
        response = client.post(
            f"/api/v1/assessments/{setup['ct1'].id}/import",
            files={"file": ("ct1.csv", csv.encode())},
            headers=auth_headers(faculty),
        )
        preview = response.json()
        assert preview["batch"]["file_format"] == "long"
        assert preview["batch"]["assessment_id"] == str(setup["ct1"].id)
        assert preview["summary"]["errors"] == 0, preview
        assert preview["summary"]["missing_students"] == 1
        client.post(
            f"/api/v1/imports/{preview['batch']['id']}/confirm", headers=auth_headers(faculty)
        )
        assert result_count(db_session) == 2

    def test_single_assessment_score_column(self, client: TestClient, faculty: User, setup) -> None:
        rows = [["Reg No", "Marks"]] + [[st.register_number, 10] for st in setup["students"]]
        response = client.post(
            f"/api/v1/assessments/{setup['ct2'].id}/import",
            files={"file": ("ct2.xlsx", xlsx(rows))},
            headers=auth_headers(faculty),
        ).json()
        marks = next(c for c in response["columns"] if c["header"] == "Marks")
        assert marks["assessment_id"] == str(setup["ct2"].id)
        assert response["summary"]["will_create"] == 3

    @pytest.mark.parametrize(
        ("name", "content", "code"),
        [
            ("m.pdf", b"%PDF-1.4", "unsupported_file"),
            ("m.csv", b"Name,CT1\nAsha,40\n", "import_rejected"),
            ("m.csv", b"Reg No,CT1,CT1\nRA1,40,41\n", "import_rejected"),
            ("m.csv", b"Reg No,CT1\n", "import_rejected"),
            ("m.csv", b"", "business_rule_violation"),
        ],
    )
    def test_rejected_files(
        self, client: TestClient, faculty: User, setup, name, content, code, db_session: Session
    ) -> None:
        response = upload(client, faculty, setup, content, name=name)
        assert response.status_code == 422
        assert response.json()["error"]["code"] == code
        assert db_session.scalar(select(func.count()).select_from(ImportBatch)) == 0


class TestAccess:
    def test_out_of_scope_upload_and_batch(
        self, client: TestClient, faculty: User, make_user, setup
    ) -> None:
        stranger = make_user(Role.FACULTY, email="stranger@srmist.edu.in")
        assert upload(client, stranger, setup, sheet(setup, [1, 2, 3])).status_code == 404
        preview = upload(client, faculty, setup, sheet(setup, [1, 2, 3])).json()
        url = f"/api/v1/imports/{preview['batch']['id']}"
        assert client.get(url, headers=auth_headers(stranger)).status_code == 404
        assert client.post(f"{url}/confirm", headers=auth_headers(stranger)).status_code == 404

    def test_only_uploader_or_administrator_edits(
        self, client: TestClient, faculty: User, make_user, hod: User, setup, db_session: Session
    ) -> None:
        from app.modules.organization.models import OfferingFaculty

        co_teacher = make_user(Role.FACULTY, email="co.teacher@srmist.edu.in")
        db_session.add(OfferingFaculty(offering_id=setup["offering"].id, user_id=co_teacher.id))
        db_session.flush()
        preview = upload(client, faculty, setup, sheet(setup, [1, 2, 3])).json()
        url = f"/api/v1/imports/{preview['batch']['id']}"
        assert client.get(f"{url}/preview", headers=auth_headers(co_teacher)).status_code == 200
        assert client.post(f"{url}/confirm", headers=auth_headers(co_teacher)).status_code == 403
        assert client.post(f"{url}/confirm", headers=auth_headers(hod)).status_code == 200

    def test_expired_batch(
        self, client: TestClient, faculty: User, setup, db_session: Session
    ) -> None:
        preview = upload(client, faculty, setup, sheet(setup, [1, 2, 3])).json()
        batch = db_session.get(ImportBatch, __import__("uuid").UUID(preview["batch"]["id"]))
        batch.expires_at = datetime.now(UTC) - timedelta(minutes=1)
        db_session.flush()
        url = f"/api/v1/imports/{batch.id}"
        assert client.post(f"{url}/confirm", headers=auth_headers(faculty)).status_code == 409
        discarded = client.post(f"{url}/discard", headers=auth_headers(faculty))
        assert discarded.status_code == 200 and discarded.json()["status"] == "discarded"

    def test_history_is_scoped(
        self, client: TestClient, faculty: User, make_user, admin: User, setup
    ) -> None:
        upload(client, faculty, setup, sheet(setup, [1, 2, 3]))
        stranger = make_user(Role.FACULTY, email="stranger@srmist.edu.in")
        assert client.get("/api/v1/imports", headers=auth_headers(faculty)).json()["total"] == 1
        assert client.get("/api/v1/imports", headers=auth_headers(stranger)).json()["total"] == 0
        assert (
            client.get("/api/v1/imports/history", headers=auth_headers(admin)).json()["total"] == 1
        )
        filtered = client.get(
            "/api/v1/imports", params={"status": "committed"}, headers=auth_headers(admin)
        ).json()
        assert filtered["total"] == 0

    def test_preview_filters(self, client: TestClient, faculty: User, setup) -> None:
        preview = upload(client, faculty, setup, sheet(setup, [45, 99, None])).json()
        url = f"/api/v1/imports/{preview['batch']['id']}/preview"
        errors = client.get(url, params={"only": "errors"}, headers=auth_headers(faculty)).json()
        issues = client.get(url, params={"only": "issues"}, headers=auth_headers(faculty)).json()
        assert [r["row"] for r in errors["rows"]] == [3]
        assert [r["row"] for r in issues["rows"]] == [3, 4]


def test_template_round_trip(client: TestClient, faculty: User, setup) -> None:
    response = client.get(
        f"/api/v1/offerings/{setup['offering'].id}/imports/template", headers=auth_headers(faculty)
    )
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    rows = list(workbook["Marks"].iter_rows(values_only=True))
    assert rows[0] == ("Register No", "Name", "CT1 (max 50)", "CT2 (max 50)")
    assert len(rows) == 4

    # Filling the template and uploading it maps every column automatically.
    filled = [list(rows[0])] + [[r[0], r[1], 40, "AB"] for r in rows[1:]]
    preview = upload(client, faculty, setup, filled).json()
    assert preview["summary"]["errors"] == 0 and preview["summary"]["will_create"] == 6

    single = client.get(
        f"/api/v1/assessments/{setup['ct1'].id}/import/template", headers=auth_headers(faculty)
    )
    assert load_workbook(io.BytesIO(single.content))["Marks"]["C1"].value == "CT1 (max 50)"
