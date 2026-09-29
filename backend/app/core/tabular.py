"""Reading uploaded CSV / Excel files into rows, with file-type detection and limits.

Used by the student bulk import and the assessment-results import pipeline. Values are
returned as trimmed strings (or None for blank cells) so each importer applies its own
typed validation; nothing is coerced to a number or zero here.
"""

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from enum import StrEnum

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.errors import BusinessRuleError

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 200


class FileType(StrEnum):
    CSV = "csv"
    XLSX = "xlsx"


class UnsupportedFileError(BusinessRuleError):
    code = "unsupported_file"


@dataclass
class Table:
    file_type: FileType
    headers: list[str]  # as found in the file (trimmed)
    rows: list[dict[str, str | None]]  # keyed by header; blank cells -> None
    row_numbers: list[int]  # 1-based spreadsheet row number of each row
    skipped_empty_rows: list[int] = field(default_factory=list)
    sheet_name: str | None = None


def detect_file_type(filename: str | None, content: bytes) -> FileType:
    """Trust the bytes, not the name: xlsx is a ZIP container; csv must be decodable text."""
    name = (filename or "").lower()
    if content[:4] == b"PK\x03\x04":
        if name.endswith((".xlsx", ".xlsm")) or not name:
            return FileType.XLSX
        raise UnsupportedFileError(
            f"'{filename}' looks like a spreadsheet/zip but is not an .xlsx file."
        )
    if content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise UnsupportedFileError("Legacy .xls files are not supported; save as .xlsx or .csv.")
    if name.endswith((".csv", ".txt")) or not name:
        try:
            _decode(content)
        except UnicodeDecodeError as exc:
            raise UnsupportedFileError("The CSV file is not valid UTF-8 text.") from exc
        return FileType.CSV
    raise UnsupportedFileError(
        f"Unsupported file type for '{filename}'. Upload a .csv or .xlsx file."
    )


def normalise_header(header: str) -> str:
    """'Reg. No ' -> 'reg_no', 'CT-1 (50)' -> 'ct_1_50'."""
    return re.sub(r"[^a-z0-9]+", "_", header.strip().lower()).strip("_")


def read_table(filename: str | None, content: bytes) -> Table:
    if not content:
        raise BusinessRuleError("The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise BusinessRuleError(f"File is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
    file_type = detect_file_type(filename, content)
    raw_rows, sheet = _read_csv(content) if file_type is FileType.CSV else _read_xlsx(content)
    return _to_table(file_type, raw_rows, sheet)


def _decode(content: bytes) -> str:
    return content.decode("utf-8-sig")


def _read_csv(content: bytes) -> tuple[list[list[str | None]], None]:
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)], None


def _read_xlsx(content: bytes) -> tuple[list[list[str | None]], str]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (
        InvalidFileException,
        KeyError,
        OSError,
        ValueError,
        # A truncated or damaged .xlsx still carries the ZIP magic, so detect_file_type accepts
        # it and openpyxl raises BadZipFile here. It derives straight from Exception -- not from
        # OSError or ValueError -- so without naming it an interrupted upload escapes as a 500
        # instead of the 422 every other unreadable file gets.
        zipfile.BadZipFile,
    ) as exc:
        raise UnsupportedFileError("The file could not be read as an .xlsx workbook.") from exc
    sheet = workbook.worksheets[0]
    rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return rows, sheet.title


def _cell(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value)


def _to_table(file_type: FileType, raw_rows: list[list[str | None]], sheet: str | None) -> Table:
    # First non-empty row is the header row.
    index = 0
    while index < len(raw_rows) and all(_blank(c) for c in raw_rows[index]):
        index += 1
    if index == len(raw_rows):
        raise BusinessRuleError("The file has no header row.")
    header_cells = raw_rows[index]
    while header_cells and _blank(header_cells[-1]):
        header_cells = header_cells[:-1]
    headers = [(c or "").strip() for c in header_cells]
    if len(headers) > MAX_COLUMNS:
        raise BusinessRuleError(f"Too many columns (maximum {MAX_COLUMNS}).")

    rows: list[dict[str, str | None]] = []
    numbers: list[int] = []
    skipped: list[int] = []
    for offset, raw in enumerate(raw_rows[index + 1 :], start=index + 2):
        cells = [(c.strip() if isinstance(c, str) else c) for c in raw]
        if all(_blank(c) for c in cells):
            skipped.append(offset)
            continue
        extra = cells[len(headers) :]
        if any(not _blank(c) for c in extra):
            raise BusinessRuleError(
                f"Row {offset} has values in columns without a header.",
                details=[{"loc": ["row", offset], "message": "Value outside header columns."}],
            )
        values = cells[: len(headers)] + [None] * (len(headers) - len(cells))
        rows.append({h: (None if _blank(v) else v) for h, v in zip(headers, values, strict=True)})
        numbers.append(offset)
        if len(rows) > MAX_ROWS:
            raise BusinessRuleError(f"Too many rows (maximum {MAX_ROWS}).")
    # Trailing blank rows are not worth reporting.
    last_data = numbers[-1] if numbers else index + 1
    skipped = [n for n in skipped if n < last_data]
    return Table(file_type, headers, rows, numbers, skipped, sheet)


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")
